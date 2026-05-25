"""生成 susr 顧問簡報 — 亮色、大字、全中文，20 張投影片。

執行：
    pip install python-pptx
    python docs/demo/gen_susr_pptx.py

輸出：docs/demo/susr-consultant-intro.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


# 品牌色（亮色主題）
BRAND_PRIMARY = RGBColor(0x1F, 0x5F, 0x4A)   # susr 深綠
BRAND_ACCENT = RGBColor(0xCC, 0x66, 0x33)    # 強調橘
BRAND_LIGHT = RGBColor(0xE8, 0xF0, 0xEB)     # 淺綠背景
TEXT_DARK = RGBColor(0x22, 0x22, 0x22)
TEXT_MUTED = RGBColor(0x55, 0x55, 0x55)
BG = RGBColor(0xFF, 0xFF, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微軟正黑體"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height


def blank_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return s


def add_text(slide, left, top, width, height, text, size=24, bold=False,
             color=TEXT_DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def add_bullets(slide, left, top, width, height, items, size=22, color=TEXT_DARK):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        r = p.add_run()
        prefix_skip = item.startswith("•") or item.startswith("  ") or item == ""
        r.text = ("• " + item) if (not prefix_skip and item) else item
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return tb


def add_title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_PRIMARY
    bar.line.fill.background()
    add_text(slide, Inches(0.4), Inches(0.1), W - Inches(0.8), Inches(0.4),
             title, size=22, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.4), Inches(0.7), W - Inches(0.8), Inches(0.5),
                 subtitle, size=18, color=TEXT_MUTED, anchor=MSO_ANCHOR.TOP)


# ========== S1 封面 ==========
s = blank_slide()
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(2.2))
bar.fill.solid()
bar.fill.fore_color.rgb = BRAND_PRIMARY
bar.line.fill.background()
add_text(s, Inches(1), Inches(0.7), W - Inches(2), Inches(0.7),
         "susr", size=72, bold=True, color=WHITE)
add_text(s, Inches(1), Inches(1.55), W - Inches(2), Inches(0.5),
         "Sustainability Brain for ESG Consultants", size=24,
         color=BRAND_LIGHT)
add_text(s, Inches(1), Inches(2.8), W - Inches(2), Inches(0.8),
         "顧問的 ESG 第二大腦", size=44, bold=True, color=TEXT_DARK)
add_text(s, Inches(1), Inches(3.8), W - Inches(2), Inches(0.6),
         "把 40-80 小時的永續報告書，壓縮到 8-15 小時",
         size=26, color=TEXT_MUTED)
add_text(s, Inches(1), Inches(6.4), W - Inches(2), Inches(0.5),
         "對象：ESG 顧問 / 永續顧問公司   ·   M4ORE  ·  2026",
         size=18, color=TEXT_MUTED)


# ========== S2 顧問痛點 ==========
s = blank_slide()
add_title_bar(s, "顧問的真實痛點：一份報告書 40-80 小時")
add_text(s, Inches(0.6), Inches(1.5), W - Inches(1.2), Inches(0.6),
         "時間都花在哪？", size=28, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.3), W - Inches(1.6), Inches(4.5), [
    "資料整理（ERP / Excel / PDF / 問卷）：30%",
    "框架對應（GRI / ISSB / 金管會）：20%",
    "章節撰寫：25%",
    "確信前置 / 證據鏈整理：15%",
    "雜事與來回：10%",
], size=24)
add_text(s, Inches(0.6), Inches(6.6), W - Inches(1.2), Inches(0.5),
         "→ 機械重複工作佔 60%+，susr 幫你拿回這段時間",
         size=20, bold=True, color=BRAND_ACCENT)


# ========== S3 Co-pilot vs Replacement ==========
s = blank_slide()
add_title_bar(s, "為什麼不是「自動報告書機器」")
add_text(s, Inches(0.6), Inches(1.4), W - Inches(1.2), Inches(0.6),
         "Co-pilot ≠ Replacement（核心戰略選擇）",
         size=28, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.2), W / 2 - Inches(0.5), Inches(4.5), [
    "AI 適合做的：",
    "  ・ 機械重複",
    "  ・ 結構化查找",
    "  ・ 框架對應檢核",
    "  ・ 證據鏈整理",
], size=22)
add_bullets(s, W / 2 + Inches(0.2), Inches(2.2), W / 2 - Inches(0.5), Inches(4.5), [
    "顧問適合做的：",
    "  ・ 雙重重大性判斷",
    "  ・ 議合策略",
    "  ・ 敘事張力",
    "  ・ 簽核與責任",
], size=22)


# ========== S4 定位 ==========
s = blank_slide()
add_title_bar(s, "susr 的定位（一句話）")
add_text(s, Inches(1), Inches(2.5), W - Inches(2), Inches(2.5),
         "「讓顧問能在同樣時間內，產出品質更高、框架更完整、"
         "故事更動人的永續報告書，責任 100% 在顧問判斷。」",
         size=32, bold=True, color=BRAND_PRIMARY, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(1), Inches(6.4), W - Inches(2), Inches(0.5),
         "— susr 三原則文件（docs/defeinition.md）",
         size=18, color=TEXT_MUTED, align=PP_ALIGN.RIGHT)


# ========== S5 給誰用 ==========
s = blank_slide()
add_title_bar(s, "給誰用 / 不給誰用")
add_text(s, Inches(0.6), Inches(1.5), W / 2 - Inches(0.8), Inches(0.6),
         "✓ 適合", size=28, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.3), W / 2 - Inches(1), Inches(4.5), [
    "獨立顧問 / 小型顧問公司",
    "  （決策快、痛點具體）",
    "服務多家客戶",
    "  （需要跨案累積方法論）",
    "重視 audit trail / 反綠色洗白",
], size=22)
add_text(s, W / 2 + Inches(0.2), Inches(1.5), W / 2 - Inches(0.8), Inches(0.6),
         "✗ 不適合", size=28, bold=True, color=BRAND_ACCENT)
add_bullets(s, W / 2 + Inches(0.4), Inches(2.3), W / 2 - Inches(1), Inches(4.5), [
    "想省顧問費的企業",
    "  （系統不取代你）",
    "只做一次性報告、無 KB 累積需求",
    "  （ROI 不划算）",
], size=22)


# ========== S6 系統概覽 ==========
s = blank_slide()
add_title_bar(s, "系統概覽 — 三層架構 + 顧問專屬功能")


def layer_box(left, top, label, items, fill_color):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             left, top, W - Inches(1.2), Inches(1.2))
    box.fill.solid()
    box.fill.fore_color.rgb = fill_color
    box.line.color.rgb = BRAND_PRIMARY
    box.line.width = Pt(1.5)
    add_text(s, left + Inches(0.3), top + Inches(0.15), Inches(3), Inches(0.5),
             label, size=22, bold=True, color=BRAND_PRIMARY)
    add_text(s, left + Inches(3.5), top + Inches(0.15), W - Inches(5), Inches(0.9),
             items, size=18, color=TEXT_DARK)


layer_box(Inches(0.6), Inches(1.3), "輸入層",
          "Excel / PDF / ERP / 問卷 / 舊報告 → RAG 自動抽指標",
          BRAND_LIGHT)
layer_box(Inches(0.6), Inches(2.7), "智能層",
          "GRI / ISSB / 金管會多框架 + 雙重重大性 + IRO 連結 + 差距分析",
          RGBColor(0xDD, 0xEA, 0xE1))
layer_box(Inches(0.6), Inches(4.1), "輸出層",
          "可編輯 Word/GDocs / 視覺化 / 版本管理 / 一鍵注入專業觀點",
          BRAND_LIGHT)
layer_box(Inches(0.6), Inches(5.5), "顧問專屬",
          "顧問 KB（學會「我們的語氣」）+ 多客戶儀表板 + 團隊協作",
          RGBColor(0xFA, 0xE6, 0xD9))


# ========== S7 entities ==========
s = blank_slide()
add_title_bar(s, "17 種實體 + 19 條 typed edges — Knowledge Graph 骨架")
add_text(s, Inches(0.6), Inches(1.4), W - Inches(1.2), Inches(0.6),
         "用力麗 5364 案例：95 pages / 70 edges / 9 entity types",
         size=24, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.3), W - Inches(1.6), Inches(4.5), [
    "Topic（議題）28 個 — 雙軸評分 + materiality_tier",
    "IRO（衝擊/風險/機會）14 個 — type / category / time_horizon",
    "Action（行動方案）14 個 — budget / progress / owner",
    "KPI（指標）8 個 — formula / boundary / framework_refs",
    "Target（目標）3 個 — 基線 / 時程 / 驗證路徑（三要素）",
    "Stakeholder / Engagement / Governance / Framework / Chapter / ...",
], size=22)


# ========== S8 Phase 3 ==========
s = blank_slide()
add_title_bar(s, "Phase 3：雙重重大性評估（MVP 核心）", "4 個 MCP tools 端到端")
add_bullets(s, Inches(0.8), Inches(1.7), W - Inches(1.6), Inches(5.2), [
    "search_topics_universe（產業議題候選池）",
    "  → 紡織業 / 觀光業 / 半導體業特化議題池",
    "score_topic_dual_axis（雙軸評分）",
    "  → Impact × Financial + 顧問 explicit 覆寫 + audit trail",
    "link_topic_to_iro（IRO 連結）",
    "  → 物理 / 轉型 / 機會 三類，自動建 topic_has_iro edge",
    "generate_materiality_matrix（矩陣 + tier_overrides）",
    "  → .md + .svg + 顧問覆寫差距 audit trail",
], size=22)


# ========== S9 Phase 7 ==========
s = blank_slide()
add_title_bar(s, "Phase 7：Gap Analysis（反綠色洗白）", "4 個 MCP tools")
add_bullets(s, Inches(0.8), Inches(1.7), W - Inches(1.6), Inches(5.2), [
    "run_compliance_checklist",
    "  → 60+ 項 GRI / ISSB / 金管會合規檢查",
    "generate_gri_content_index",
    "  → 自動產出 GRI Content Index 表 + 涵蓋率",
    "generate_assurance_readiness_checklist",
    "  → 確信師上場前的 KPI / DataPoint 自我健檢",
    "run_full_gap_analysis（整合 aggregator）",
    "  → critical / warning / info 三層 severity + 修補建議",
], size=22)


# ========== S10 Phase 6 ==========
s = blank_slide()
add_title_bar(s, "Phase 6：報告組裝（payload prep）", "4 個 MCP tools")
add_bullets(s, Inches(0.8), Inches(1.7), W - Inches(1.6), Inches(5.2), [
    "prepare_docx_payload — 整份報告書（chapters + matrix + KPIs）",
    "prepare_pdf_payload — 同 + iXBRL 標記預留（金管會 2028）",
    "prepare_pptx_investor_deck — 10-15 頁投資人版",
    "prepare_pptx_board_deck — 5-8 頁董事會版",
    "",
    "重點：susr 不渲染文檔，只組裝資料",
    "  → 真實渲染由 anthropics/skills 的 docx/pdf/pptx skill 處理",
    "  → 保留所有 source reference + tier_overrides warning",
], size=22)


# ========== S11 invariants ==========
s = blank_slide()
add_title_bar(s, "6 條連結性 Invariants（反綠色洗白護城河）",
              "lealea fixture 6/6 全綠 ✓")
items = [
    ("I1a", "核心 topic 必有 IRO"),
    ("I1b", "IRO 必有 Action"),
    ("I2", "Chapter conformsTo Framework"),
    ("I3", "Target 三要素齊（基線 / 時程 / 驗證路徑）"),
    ("I4", "DataPoint 可追溯（必有 SourceDoc）"),
    ("I5", "排放 DataPoint 對應有效 EmissionFactor"),
]
for i, (code, desc) in enumerate(items):
    top = Inches(1.5 + i * 0.85)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.7), top, Inches(1.2), Inches(0.6))
    box.fill.solid()
    box.fill.fore_color.rgb = BRAND_PRIMARY
    box.line.fill.background()
    add_text(s, Inches(0.7), top, Inches(1.2), Inches(0.6),
             code, size=22, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.1), top + Inches(0.05),
             W - Inches(2.8), Inches(0.5),
             desc, size=22, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)


# ========== S12 audit trail ==========
s = blank_slide()
add_title_bar(s, "Audit Trail：page_versions + timeline",
              "確信師質疑時，一秒給整條鏈")
add_text(s, Inches(0.6), Inches(1.5), W - Inches(1.2), Inches(0.6),
         "每個 KPI 數字背後：", size=26, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.3), W - Inches(1.6), Inches(3), [
    "來源（SourceDoc / 環境部排放因子版本）",
    "公式（kWh × factor / 員工數 / 房間數 / 營收）",
    "計算邊界（合併 / 營運控制 / 股權法）",
    "責任人 / 驗證者 / 確信狀態",
    "時間戳 + 操作紀錄（ingest / verify / restate / assure）",
], size=22)
add_text(s, Inches(0.6), Inches(5.6), W - Inches(1.2), Inches(0.6),
         "跨年 restate > 5% → 自動偵測 + 強制揭露原因",
         size=22, bold=True, color=BRAND_ACCENT)
add_text(s, Inches(0.6), Inches(6.3), W - Inches(1.2), Inches(0.5),
         "反綠色洗白 hard rule，工具層強制（不能用「忘了揭」當藉口）",
         size=18, color=TEXT_MUTED)


# ========== S13 場景 1 ==========
s = blank_slide()
add_title_bar(s, "場景 1：接到新客戶第一週", "30 秒建立 workspace")
add_text(s, Inches(0.6), Inches(1.5), W - Inches(1.2), Inches(0.6),
         "顧問：", size=22, bold=True, color=BRAND_PRIMARY)
add_text(s, Inches(0.8), Inches(2.1), W - Inches(1.6), Inches(1),
         "「幫我建一個叫 acme-1234 的客戶 workspace。"
         "紡織業上市櫃，2025 永續報告書，營運控制法，GRI + ISSB + 金管會。」",
         size=22, color=TEXT_DARK)
add_text(s, Inches(0.6), Inches(3.5), W - Inches(1.2), Inches(0.6),
         "Claude Desktop（透過 susr-mcp）：",
         size=22, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(4.1), W - Inches(1.6), Inches(3), [
    "建立 ~/susr-clients/acme-1234/（獨立 git repo）",
    "初始 commit + shared_kb snapshot 13 個檔案",
    "entities/{topics,stakeholders,governance,kpis,targets}/ 結構",
    "projects/2025-sustainability-report/ + .susr/db.sqlite",
], size=20)


# ========== S14 場景 2 ==========
s = blank_slide()
add_title_bar(s, "場景 2：跨年延續（第二年怎麼用上一年資料）")
add_bullets(s, Inches(0.8), Inches(1.6), W - Inches(1.6), Inches(5.5), [
    "cp -r projects/2025-sr/  projects/2026-sr/",
    "  → fork 為 2026 baseline（含 framework_version）",
    "entities/ 跨年常駐（議題池、利害關係人、KPI 定義、目標）",
    "  → 不必重評議題（只增量更新）",
    "  → KPI 自動 carry over values_by_year",
    "Restatement 跨年比對：> 5% 自動偵測 + 揭露",
    "  → 反綠色洗白機制不請假",
    "",
    "→ 第二年顧問時間從 40-80h 進一步壓縮到 4-8h",
], size=22)


# ========== S15 場景 3 ==========
s = blank_slide()
add_title_bar(s, "場景 3：客戶問「為什麼用這個排放因子？」")
add_text(s, Inches(0.6), Inches(1.4), W - Inches(1.2), Inches(0.6),
         "確信師 / 客戶質疑時的證據鏈追溯",
         size=26, bold=True, color=BRAND_PRIMARY)
add_text(s, Inches(0.8), Inches(2.3), W - Inches(1.6), Inches(0.6),
         "顧問：「acme-1234 Scope 2 用什麼排放因子？」",
         size=22, color=TEXT_DARK)
add_text(s, Inches(0.6), Inches(3.3), W - Inches(1.2), Inches(0.6),
         "Brain 一秒回應：", size=22, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(3.9), W - Inches(1.6), Inches(3), [
    "EmissionFactor：環境部 113 年度（version=2024-Q1）",
    "effective_from: 2024-04-01 ~ effective_to: 2025-03-31",
    "由顧問 Mary 於 2024-04-15 ingest，verified by John",
    "DataPoint 經 derivedFrom 連到台電帳單 PDF",
    "確信狀態：第三方有限確信（assurance: limited）",
], size=20)


# ========== S16 場景 4 ==========
s = blank_slide()
add_title_bar(s, "場景 4：把客戶經驗變方法論（顧問 KB 單向闘）")
add_text(s, Inches(0.6), Inches(1.4), W - Inches(1.2), Inches(0.6),
         "promote_to_consultant_kb 兩步驟 confirm 機制",
         size=24, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.2), W - Inches(1.6), Inches(2.5), [
    "Step 1: dry-run → 回 {needs_confirm: True, confirm_token: <sha>}",
    "Step 2: 帶 token + reviewer 簽核 → 真寫 consultant-kb + audit log",
], size=20)
add_text(s, Inches(0.6), Inches(5), W - Inches(1.2), Inches(0.6),
         "Hard Rules（工具層強制）：",
         size=22, bold=True, color=BRAND_ACCENT)
add_bullets(s, Inches(0.8), Inches(5.6), W - Inches(1.6), Inches(1.5), [
    "Client 之間絕不互通（每個 client 獨立 git repo）",
    "Client → consultant-kb 必須手動匿名化 + reviewer 審核",
    "consultant-kb → 新 client 可讀（方法論複用）",
], size=20)


# ========== S17 安裝 ==========
s = blank_slide()
add_title_bar(s, "安裝與啟動", "5 分鐘從零到能用")
add_text(s, Inches(0.6), Inches(1.4), W - Inches(1.2), Inches(0.6),
         "1. 安裝", size=24, bold=True, color=BRAND_PRIMARY)
add_text(s, Inches(0.8), Inches(2), W - Inches(1.6), Inches(0.5),
         "pip install susr", size=22, color=TEXT_DARK)
add_text(s, Inches(0.6), Inches(3), W - Inches(1.2), Inches(0.6),
         "2. 設定 Claude Desktop MCP config",
         size=24, bold=True, color=BRAND_PRIMARY)
config_text = (
    '{\n  "mcpServers": {\n    "susr": { "command": "susr-mcp" }\n  }\n}'
)
add_text(s, Inches(0.8), Inches(3.6), W - Inches(1.6), Inches(1.5),
         config_text, size=18, color=TEXT_DARK)
add_text(s, Inches(0.6), Inches(5.3), W - Inches(1.2), Inches(0.6),
         "3. 重啟 Claude Desktop → 開新對話",
         size=24, bold=True, color=BRAND_PRIMARY)
add_text(s, Inches(0.8), Inches(5.9), W - Inches(1.6), Inches(0.5),
         "右下角 icon 含 19 個 susr 工具就 OK，開始對話",
         size=20, color=TEXT_DARK)


# ========== S18 安全隱私 ==========
s = blank_slide()
add_title_bar(s, "安全與隱私")
add_bullets(s, Inches(0.8), Inches(1.6), W - Inches(1.6), Inches(5.5), [
    "Local-first — 所有資料在你本機（不上雲）",
    "  brain DB = .susr/db.sqlite，可備份 / 加密 / 移動",
    "Per-client 獨立 git repo",
    "  可推到 client 自有 GitHub / GitLab / private remote",
    "顧問 KB 單向闘（工具層強制 hard rule）",
    "  client 之間絕不互通；KB ← client 必手動匿名 + review",
    "預設本地 embedding（BGE-M3）",
    "  client-level 可選雲端（顧問需在客戶授權後切換）",
], size=22)


# ========== S19 限制揭露 ==========
s = blank_slide()
add_title_bar(s, "限制（誠實揭露）", "對使用者誠實的前提，是先對自己誠實")
add_bullets(s, Inches(0.8), Inches(1.6), W - Inches(1.6), Inches(5.5), [
    "✗ Phase 4 GHG 計算引擎未實作",
    "  顧問仍用 xlsx skill 手算 Scope 1/2/3",
    "✗ Phase 5 章節 LLM 自動草稿未實作",
    "  顧問用 brain 整理素材，章節敘事仍人寫",
    "✗ Phase 8 公告 / MOPS filing 未實作",
    "⏳ Claude Desktop stdio 端到端待 mcp SDK 環境驗證",
    "",
    "→ susr v0.1 = Phase 3 重大性評估 + Phase 6 payload + Phase 7 gap",
    "  其他 phases 仍仰賴顧問既有工具與判斷",
], size=22)


# ========== S20 加入方式 ==========
s = blank_slide()
add_title_bar(s, "怎麼加入 / 試用 / 給回饋")
add_text(s, Inches(0.6), Inches(1.5), W - Inches(1.2), Inches(0.6),
         "現在可以做的：", size=26, bold=True, color=BRAND_PRIMARY)
add_bullets(s, Inches(0.8), Inches(2.2), W - Inches(1.6), Inches(2.5), [
    "GitHub: M4ORE/m4ore-skill-susr",
    "pip install susr",
    "閱讀 docs/user-guide/scenarios/（4 個顧問場景）",
    "閱讀 docs/demo/lealea-phase3-quickstart.md（30 分鐘逐句腳本）",
], size=22)
add_text(s, Inches(0.6), Inches(4.9), W - Inches(1.2), Inches(0.6),
         "我們在找 design partner：",
         size=26, bold=True, color=BRAND_ACCENT)
add_bullets(s, Inches(0.8), Inches(5.6), W - Inches(1.6), Inches(1.5), [
    "3-5 家獨立顧問 / 小型顧問公司",
    "願意提供 1-2 個客戶 case 做 design partnership",
    "回饋頻率高、共建方法論",
], size=22)


# Save
out = Path(__file__).resolve().parent / "susr-consultant-intro.pptx"
prs.save(str(out))
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
print(f"Size: {out.stat().st_size:,} bytes")
